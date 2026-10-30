import csv
import io

import openpyxl
from django.core.exceptions import PermissionDenied
from django.http import Http404, HttpResponse
from django.shortcuts import render
from django.utils import timezone
from openpyxl.styles import Font, PatternFill

from apps.audit import services as audit
from django.contrib.auth.decorators import login_required
from apps.employees.models import Department

from .reports import REPORTS, allowed, scoped_employees


def _safe(v):
    """Neutralise spreadsheet formula injection in exported text cells."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


@login_required
def index(request):
    groups = {"employee": [], "manager": [], "admin": []}
    for r in REPORTS.values():
        if allowed(request.user, r):
            groups[r.audience].append(r)
    return render(request, "reporting/index.html", {"groups": groups})


@login_required
def report(request, key):
    rep = REPORTS.get(key)
    if rep is None:
        raise Http404
    if not allowed(request.user, rep):
        raise PermissionDenied
    f = request.GET
    emps = scoped_employees(request.user, f, rep.audience)
    cols, rows = rep.fn(request.user, f, emps)
    fmt = f.get("format")
    if fmt in ("csv", "xlsx"):
        audit.log(request.user, "report.exported", None, target_type="report", target_id=key, after={"format": fmt, "rows": len(rows),
                                                                                                    "filters": {k: v for k, v in f.items() if k != "format"}})
        stamp = timezone.localdate().strftime("%Y%m%d")
        if fmt == "csv":
            resp = HttpResponse(content_type="text/csv; charset=utf-8")
            resp["Content-Disposition"] = f'attachment; filename="{key}_{stamp}.csv"'
            resp.write("﻿")
            w = csv.writer(resp)
            w.writerow(cols)
            for row in rows:
                w.writerow([_safe(str(c)) if c is not None else "" for c in row])
            return resp
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = rep.title[:31]
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
        for row in rows:
            ws.append([_safe(c) if not hasattr(c, "normalize") else float(c) for c in row])
        for i, _ in enumerate(cols, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="{key}_{stamp}.xlsx"'
        return resp
    return render(request, "reporting/report.html", {"rep": rep, "cols": cols, "rows": rows[:500], "total": len(rows),
                                                     "departments": Department.objects.filter(is_active=True), "today": timezone.localdate()})
