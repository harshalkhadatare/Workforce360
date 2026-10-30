"""Attendance Excel import pipeline: upload → map → validate/preview → confirm → import → results.

Reconciliation rules (documented in README "Attendance import"):
  * Duplicate key = Employee ID | punch timestamp (to the second) | punch type.
    A row whose punches all already exist (live) is DUPLICATE and skipped.
  * Multiple punches per employee per day are allowed (different timestamps).
  * DAILY format only: if an earlier batch already holds a different In (or Out)
    time for the same employee/date, the row is a CONFLICT. It is NOT imported
    automatically; an administrator must accept (creates an audited adjustment,
    originals kept) or reject it on the results page.
  * Import runs in ONE database transaction. If anything fails, nothing is
    imported and the batch is marked FAILED with the error.
"""
import datetime
import io
import logging

import openpyxl
import pandas as pd
from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Count
from django.utils import timezone

from apps.attendance.models import AttendanceAdjustment, AttendancePunch
from apps.attendance.services import recalculate_month
from apps.audit import services as audit
from apps.core.validators import sha256_of, validate_upload
from apps.employees.models import Employee

from . import parsing as P
from .models import AttendanceImportBatch, AttendanceImportRow, ImportFormat

log = logging.getLogger(__name__)
RS = AttendanceImportRow.RowStatus
BS = AttendanceImportBatch.Status


def tz():
    return timezone.get_current_timezone()


def aware(naive):
    return timezone.make_aware(naive, tz())


# ---------------------------------------------------------------------------
# Step 1 — upload
# ---------------------------------------------------------------------------
def create_batch(upload, actor):
    validate_upload(upload, allowed=(".xlsx",), max_bytes=settings.MAX_IMPORT_BYTES)
    try:
        upload.seek(0)
        wb = openpyxl.load_workbook(upload, read_only=True, data_only=True)
        sheets = list(wb.sheetnames)
        wb.close()
    except Exception as exc:  # noqa: BLE001
        raise ValidationError(f"The workbook could not be read: {exc.__class__.__name__}. Save it as .xlsx and try again.")
    if not sheets:
        raise ValidationError("The workbook contains no worksheets.")
    digest = sha256_of(upload)
    previous = AttendanceImportBatch.objects.filter(sha256=digest, status__in=[BS.COMPLETED, BS.COMPLETED_WITH_ERRORS]).first()
    batch = AttendanceImportBatch(original_name=upload.name[:200], sha256=digest, size_bytes=upload.size, sheet_names=sheets,
                                  uploaded_by=actor, same_file_batch=previous, sheet_name=sheets[0])
    batch.file.save(upload.name, upload, save=False)
    batch.save()
    audit.log(actor, "import.uploaded", batch, after={"file": batch.original_name, "sha256": digest})
    return batch


# ---------------------------------------------------------------------------
# Step 2 — identify format
# ---------------------------------------------------------------------------
def read_frame(batch, nrows=None):
    with batch.file.open("rb") as fh:
        data = fh.read()
    try:
        df = pd.read_excel(io.BytesIO(data), sheet_name=batch.sheet_name, header=batch.header_row - 1, dtype=object,
                           engine="openpyxl", nrows=nrows)
    except ValueError as exc:
        raise ValidationError(f"Could not read worksheet '{batch.sheet_name}': {exc}")
    df.columns = [str(c).strip() for c in df.columns]
    df = df.dropna(how="all")
    return df


def detect_columns(batch):
    df = read_frame(batch, nrows=50)
    cols = [c for c in df.columns if not c.startswith("Unnamed:")]
    batch.detected_columns = cols
    batch.save(update_fields=["detected_columns", "updated_at"])
    return cols, df.head(10)


ALIASES = {
    "employee_id": ["employee id", "emp id", "empid", "emp code", "employee code", "ecode", "staff id", "card no", "user id"],
    "employee_name": ["employee name", "name", "emp name"],
    "attendance_date": ["attendance date", "date", "att date", "punch date"],
    "in_time": ["in time", "intime", "in", "first in", "check in", "in punch"],
    "out_time": ["out time", "outtime", "out", "last out", "check out", "out punch"],
    "punch_timestamp": ["punch timestamp", "punch time", "timestamp", "log time", "datetime", "punch datetime"],
    "punch_type": ["punch type", "direction", "in/out", "io type", "type"],
    "device": ["device", "device name", "location", "terminal", "reader"],
    "status": ["status", "attendance status"],
}


def suggest_mapping(columns):
    lower = {c.lower().strip(): c for c in columns}
    out = {}
    for field, names in ALIASES.items():
        for n in names:
            if n in lower:
                out[field] = lower[n]
                break
    return out


def validate_mapping(format_mode, mapping, columns):
    missing = [c for c in mapping.values() if c and c not in columns]
    if missing:
        raise ValidationError(f"Mapped column(s) not found in the sheet: {', '.join(missing)}")
    if not mapping.get("employee_id"):
        raise ValidationError("Map the Employee ID column.")
    if format_mode == ImportFormat.DAILY:
        if not mapping.get("attendance_date"):
            raise ValidationError("Daily format requires the Attendance date column.")
        if not (mapping.get("in_time") or mapping.get("out_time")):
            raise ValidationError("Daily format requires an In time and/or Out time column.")
    elif format_mode == ImportFormat.PUNCH_LOG:
        if not (mapping.get("punch_timestamp") or (mapping.get("attendance_date") and mapping.get("in_time"))):
            raise ValidationError("Punch-log format requires a Punch timestamp column (or Date + Time columns).")
    else:
        raise ValidationError("Choose the file layout.")


# ---------------------------------------------------------------------------
# Steps 3–4 — preview & validation
# ---------------------------------------------------------------------------
class _Ctx:
    def __init__(self, batch):
        self.batch = batch
        self.today = timezone.localdate()
        self.employees = {e.employee_id.upper(): e for e in Employee.objects.all()}
        self.seen_keys = {}
        self.live_keys = {}

    def load_live_keys(self, keys):
        missing = [k for k in keys if k not in self.live_keys]
        if missing:
            found = dict(AttendancePunch.objects.filter(source_key__in=missing, is_void=False).values_list("source_key", "batch_id"))
            for k in missing:
                self.live_keys[k] = found.get(k)


def _get(raw, mapping, field):
    col = mapping.get(field)
    return raw.get(col) if col else None


def evaluate_row(ctx, raw, row_number):
    """Parse and validate one row. Returns an unsaved AttendanceImportRow."""
    b = ctx.batch
    m = b.mapping
    row = AttendanceImportRow(batch=b, row_number=row_number, raw_data={k: P.jsonable(v) for k, v in raw.items()})
    msgs = []

    def err(t):
        msgs.append({"level": "error", "text": t})

    def warn(t):
        msgs.append({"level": "warning", "text": t})

    code = P.parse_employee_id(_get(raw, m, "employee_id"))
    row.employee_code = code
    emp = ctx.employees.get(code) if code else None
    if not code:
        err("Employee ID is missing")
    elif emp is None:
        err(f"Unknown employee ID '{code}'")
    elif not emp.is_active:
        err(f"Employee {code} is deactivated")
    row.employee = emp

    punches = []  # (naive_dt, type)
    if b.format_mode == ImportFormat.DAILY:
        try:
            d = P.parse_date(_get(raw, m, "attendance_date"), b.dayfirst)
            row.attendance_date = d
        except P.ParseError as e:
            err(str(e))
            d = None
        tin = tout = None
        try:
            tin = P.parse_time(_get(raw, m, "in_time"))
        except P.ParseError as e:
            err(f"In time: {e}")
        try:
            tout = P.parse_time(_get(raw, m, "out_time"))
        except P.ParseError as e:
            err(f"Out time: {e}")
        if d:
            if tin is None and tout is None and not any(x["level"] == "error" and "time" in x["text"] for x in msgs):
                err("Both In and Out times are empty")
            if tin and tout and tout <= tin:
                err(f"Out time {tout:%H:%M} is not after In time {tin:%H:%M} (conflicting punches)")
            if tin:
                punches.append((datetime.datetime.combine(d, tin), "IN"))
                row.in_time = aware(punches[-1][0])
            if tout:
                punches.append((datetime.datetime.combine(d, tout), "OUT"))
                row.out_time = aware(punches[-1][0])
            if tin and not tout:
                warn("Out time missing — day will show as a missing punch")
            if tout and not tin:
                warn("In time missing — day will show as a missing punch")
    else:
        try:
            if m.get("punch_timestamp"):
                ts = P.parse_timestamp(_get(raw, m, "punch_timestamp"), b.dayfirst)
            else:
                d = P.parse_date(_get(raw, m, "attendance_date"), b.dayfirst)
                t = P.parse_time(_get(raw, m, "in_time"))
                if t is None:
                    raise P.ParseError("Punch time is missing")
                ts = datetime.datetime.combine(d, t)
            row.attendance_date = ts.date()
            row.punch_time = aware(ts)
        except P.ParseError as e:
            err(str(e))
            ts = None
        ptype = "UNKNOWN"
        if m.get("punch_type"):
            try:
                ptype = P.parse_punch_type(_get(raw, m, "punch_type"))
            except P.ParseError as e:
                err(str(e))
        elif ts:
            warn("No punch type column — first and last punch of the day will be used as In and Out")
        row.punch_type = ptype
        if ts:
            punches.append((ts, ptype))

    if m.get("device"):
        row.device = str(P.jsonable(_get(raw, m, "device")) or "")[:80]
    if row.attendance_date:
        if row.attendance_date > ctx.today:
            err(f"Date {row.attendance_date:%d %b %Y} is in the future")
        if emp and row.attendance_date < emp.date_of_joining:
            err(f"Date is before {code}'s joining date ({emp.date_of_joining:%d %b %Y})")
        if emp and emp.exit_date and row.attendance_date > emp.exit_date:
            err(f"Date is after {code}'s exit date")
    name_val = _get(raw, m, "employee_name")
    if emp and not P.is_blank(name_val):
        given = " ".join(str(name_val).split()).lower()
        if given not in (emp.full_name.lower(), f"{emp.first_name} {emp.last_name}".strip().lower()):
            warn(f"Name '{name_val}' differs from the employee master ('{emp.full_name}')")

    keys = [P.business_key(code, ts, pt) for ts, pt in punches] if code else []
    row.business_keys = keys
    row.messages = msgs
    if any(x["level"] == "error" for x in msgs):
        row.status = RS.ERROR
        return row

    # Duplicates: within this file, then against live punches in the database
    ctx.load_live_keys(keys)
    in_file = [k for k in keys if k in ctx.seen_keys]
    in_db = [k for k in keys if ctx.live_keys.get(k) is not None]
    for k in keys:
        ctx.seen_keys.setdefault(k, row_number)
    if keys and len(set(in_file + in_db)) == len(keys):
        row.status = RS.DUPLICATE
        if in_file:
            msgs.append({"level": "warning", "text": f"Duplicate of row {ctx.seen_keys[in_file[0]]} in this file"})
        else:
            msgs.append({"level": "warning", "text": f"Already imported (batch IMP-{ctx.live_keys[in_db[0]]:05d})"})
        return row

    # Conflicts with an earlier import (daily layout only)
    if b.format_mode == ImportFormat.DAILY and emp and row.attendance_date:
        existing = AttendancePunch.objects.filter(employee=emp, punch_date=row.attendance_date, is_void=False).exclude(batch=b)
        ex_in = [p.punch_time for p in existing if p.punch_type == "IN"]
        ex_out = [p.punch_time for p in existing if p.punch_type == "OUT"]
        conflicts = []
        if row.in_time and ex_in and row.in_time not in ex_in:
            conflicts.append(f"In {timezone.localtime(min(ex_in)):%H:%M} → {timezone.localtime(row.in_time):%H:%M}")
        if row.out_time and ex_out and row.out_time not in ex_out:
            conflicts.append(f"Out {timezone.localtime(max(ex_out)):%H:%M} → {timezone.localtime(row.out_time):%H:%M}")
        if conflicts:
            msgs.append({"level": "warning", "text": "Differs from previously imported data: " + ", ".join(conflicts)
                         + ". Needs reconciliation."})
            row.status = RS.CONFLICT
            return row

    row.status = RS.WARNING if any(x["level"] == "warning" for x in msgs) else RS.VALID
    return row


@transaction.atomic
def validate_batch(batch, actor):
    if batch.is_final:
        raise ValidationError("This batch has already been processed.")
    validate_mapping(batch.format_mode, batch.mapping, batch.detected_columns)
    df = read_frame(batch)
    if len(df) == 0:
        raise ValidationError("The selected worksheet has no data rows below the header.")
    if len(df) > settings.MAX_IMPORT_ROWS:
        raise ValidationError(f"The worksheet has {len(df)} rows; the limit is {settings.MAX_IMPORT_ROWS}. Split the file.")
    batch.rows.all().delete()
    ctx = _Ctx(batch)
    rows = []
    for idx, raw in zip(df.index, df.to_dict(orient="records")):
        rows.append(evaluate_row(ctx, raw, int(idx) + batch.header_row + 1))
    AttendanceImportRow.objects.bulk_create(rows, batch_size=1000)
    _update_counts(batch)
    batch.status = BS.VALIDATED
    batch.save()
    audit.log(actor, "import.validated", batch, after=_counts(batch))
    return batch


def _update_counts(batch):
    counts = {s: 0 for s in RS.values}
    for st, n in batch.rows.order_by().values_list("status").annotate(n=Count("id")):
        counts[st] = n
    batch.total_rows = sum(counts.values())
    batch.valid_rows = counts[RS.VALID] + counts[RS.WARNING]
    batch.warning_rows = counts[RS.WARNING]
    batch.error_rows = counts[RS.ERROR]
    batch.duplicate_rows = counts[RS.DUPLICATE]
    batch.conflict_rows = counts[RS.CONFLICT]
    batch.imported_rows = counts[RS.IMPORTED] + counts[RS.RECON_ACCEPTED]


def _counts(batch):
    return {f: getattr(batch, f) for f in ["total_rows", "valid_rows", "warning_rows", "error_rows", "duplicate_rows",
                                           "conflict_rows", "imported_rows", "punches_created"]}


# ---------------------------------------------------------------------------
# Steps 5–6 — confirm & import
# ---------------------------------------------------------------------------
def _create_punches(row, batch):
    created = 0
    pairs = []
    if row.punch_time:
        pairs.append((row.punch_time, row.punch_type or "UNKNOWN"))
    else:
        if row.in_time:
            pairs.append((row.in_time, "IN"))
        if row.out_time:
            pairs.append((row.out_time, "OUT"))
    for (ts, ptype), key in zip(pairs, row.business_keys):
        try:
            with transaction.atomic():
                AttendancePunch.objects.create(
                    employee=row.employee, punch_time=ts, punch_date=timezone.localtime(ts).date(), punch_type=ptype,
                    device=row.device, batch=batch, import_row=row, source_key=key,
                    raw_value=f"{timezone.localtime(ts):%Y-%m-%d %H:%M:%S}",
                )
                created += 1
        except IntegrityError:
            continue  # imported concurrently by another batch — treated as duplicate
    return created


def _recalc(affected):
    for emp_id, months in affected.items():
        emp = Employee.objects.get(pk=emp_id)
        for y, mth in sorted(months):
            recalculate_month(emp, y, mth)


def import_batch(batch, actor):
    """Import all VALID/WARNING rows in one transaction. Returns the batch."""
    try:
        with transaction.atomic():
            batch = AttendanceImportBatch.objects.select_for_update().get(pk=batch.pk)
            if batch.status != BS.VALIDATED:
                raise ValidationError("Only a validated batch awaiting confirmation can be imported.")
            affected, dates, punches = {}, [], 0
            for row in batch.rows.select_related("employee").filter(status__in=[RS.VALID, RS.WARNING]):
                n = _create_punches(row, batch)
                if n == 0:
                    row.status = RS.DUPLICATE
                    row.messages = row.messages + [{"level": "warning", "text": "Skipped — punches already exist"}]
                else:
                    row.status = RS.IMPORTED
                    punches += n
                    dates.append(row.attendance_date)
                    affected.setdefault(row.employee_id, set()).add((row.attendance_date.year, row.attendance_date.month))
                row.save(update_fields=["status", "messages"])
            batch.punches_created = punches
            if dates:
                batch.date_from, batch.date_to = min(dates), max(dates)
            _update_counts(batch)
            batch.status = BS.COMPLETED if not (batch.error_rows or batch.conflict_rows) else BS.COMPLETED_WITH_ERRORS
            batch.confirmed_by, batch.confirmed_at, batch.completed_at = actor, timezone.now(), timezone.now()
            batch.save()
            _recalc(affected)
            audit.log(actor, "import.completed", batch, after=_counts(batch))
            transaction.on_commit(lambda: _notify_import(batch.pk))
            return batch
    except ValidationError:
        raise
    except Exception as exc:  # noqa: BLE001 — record the failure; the transaction rolled everything back
        log.exception("Import batch %s failed", batch.pk)
        AttendanceImportBatch.objects.filter(pk=batch.pk).update(
            status=BS.FAILED, failure_message=f"{exc.__class__.__name__}: {exc}"[:2000], completed_at=timezone.now())
        audit.log(actor, "import.failed", batch, reason=str(exc)[:500])
        batch.refresh_from_db()
        return batch


def _notify_import(batch_pk):
    from django.urls import reverse

    from apps.notifications.services import notify_users

    b = AttendanceImportBatch.objects.get(pk=batch_pk)
    if b.uploaded_by:
        notify_users([b.uploaded_by], event_type="import.completed", title=f"Attendance import {b.batch_code} {b.get_status_display().lower()}",
                     intro=f"{b.original_name}: {b.imported_rows} rows imported, {b.error_rows} rejected, {b.duplicate_rows} duplicates, "
                           f"{b.conflict_rows} need reconciliation.",
                     link_path=reverse("imports:results", args=[b.pk]), dedupe=f"import:{b.pk}:{b.status}", related=b)


def discard_batch(batch, actor):
    if batch.is_final:
        raise ValidationError("This batch can no longer be discarded.")
    batch.status = BS.DISCARDED
    batch.save(update_fields=["status", "updated_at"])
    audit.log(actor, "import.discarded", batch)


# ---------------------------------------------------------------------------
# Step 7 — results, reconciliation, retry, rollback
# ---------------------------------------------------------------------------
@transaction.atomic
def reconcile_row(row, actor, accept, note=""):
    row = AttendanceImportRow.objects.select_for_update(of=("self",)).select_related("batch", "employee").get(pk=row.pk)
    if row.status != RS.CONFLICT:
        raise ValidationError("This row does not need reconciliation.")
    if row.batch.status not in (BS.COMPLETED, BS.COMPLETED_WITH_ERRORS):
        raise ValidationError("Reconcile after the batch has been imported.")
    from apps.attendance.services import effective_punches, recalculate_for_dates

    if accept:
        _, _, _, orig_in, orig_out = effective_punches(row.employee, row.attendance_date)
        AttendanceAdjustment.objects.filter(employee=row.employee, date=row.attendance_date, is_active=True).update(is_active=False)
        AttendanceAdjustment.objects.create(
            employee=row.employee, date=row.attendance_date, original_in=orig_in, original_out=orig_out,
            corrected_in=row.in_time, corrected_out=row.out_time,
            reason=f"Corrected source data in {row.batch.batch_code} row {row.row_number}. {note}".strip(),
            source=AttendanceAdjustment.Source.IMPORT_RECONCILIATION, import_row=row, approved_by=actor,
        )
        row.status = RS.RECON_ACCEPTED
        recalculate_for_dates(row.employee, [row.attendance_date])
    else:
        row.status = RS.RECON_REJECTED
    row.reconciled_by, row.reconciled_at = actor, timezone.now()
    row.messages = row.messages + [{"level": "info", "text": f"Reconciled by {actor.employee_code}: {'accepted' if accept else 'existing data kept'}. {note}".strip()}]
    row.save()
    batch = row.batch
    _update_counts(batch)
    batch.save()
    audit.log(actor, "import.reconciled", row.batch, employee_ref=row.employee_code,
              after={"row": row.row_number, "accepted": accept, "in": row.in_time, "out": row.out_time}, reason=note)
    return row


def retry_rejected_rows(batch, actor):
    """Re-validate rejected rows (e.g. after a missing employee was created) and import the ones now valid."""
    with transaction.atomic():
        batch = AttendanceImportBatch.objects.select_for_update().get(pk=batch.pk)
        if batch.status not in (BS.COMPLETED, BS.COMPLETED_WITH_ERRORS):
            raise ValidationError("Only completed batches can retry rejected rows.")
        ctx = _Ctx(batch)
        affected, imported = {}, 0
        for row in batch.rows.filter(status=RS.ERROR):
            fresh = evaluate_row(ctx, row.raw_data, row.row_number)
            for f in ["employee", "employee_code", "attendance_date", "in_time", "out_time", "punch_time", "punch_type",
                      "device", "status", "messages", "business_keys"]:
                setattr(row, f, getattr(fresh, f))
            if row.status in (RS.VALID, RS.WARNING):
                if _create_punches(row, batch):
                    row.status = RS.IMPORTED
                    imported += 1
                    affected.setdefault(row.employee_id, set()).add((row.attendance_date.year, row.attendance_date.month))
                    batch.date_from = min(filter(None, [batch.date_from, row.attendance_date]))
                    batch.date_to = max(filter(None, [batch.date_to, row.attendance_date]))
                else:
                    row.status = RS.DUPLICATE
            row.save()
        _update_counts(batch)
        batch.status = BS.COMPLETED if not (batch.error_rows or batch.conflict_rows) else BS.COMPLETED_WITH_ERRORS
        batch.save()
        _recalc(affected)
        audit.log(actor, "import.retried", batch, after={"newly_imported": imported})
    return imported


@transaction.atomic
def rollback_batch(batch, actor, reason):
    if not (reason or "").strip():
        raise ValidationError("A reason is required to roll back an import.")
    batch = AttendanceImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.status not in (BS.COMPLETED, BS.COMPLETED_WITH_ERRORS):
        raise ValidationError("Only completed imports can be rolled back.")
    affected = {}
    for p in batch.punches.filter(is_void=False):
        affected.setdefault(p.employee_id, set()).add((p.punch_date.year, p.punch_date.month))
    n = batch.punches.filter(is_void=False).update(is_void=True, void_reason=f"Batch rolled back: {reason}"[:255])
    AttendanceAdjustment.objects.filter(import_row__batch=batch).update(is_active=False)
    batch.status = BS.ROLLED_BACK
    batch.rolled_back_by, batch.rolled_back_at, batch.rollback_reason = actor, timezone.now(), reason
    batch.save()
    _recalc(affected)
    audit.log(actor, "import.rolled_back", batch, reason=reason, after={"punches_voided": n})
    return n


def error_report_workbook(batch):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Import results"
    raw_cols = batch.detected_columns or []
    ws.append(["Row", "Result", "Messages"] + raw_cols)
    for row in batch.rows.exclude(status__in=[RS.IMPORTED, RS.VALID]).order_by("row_number"):
        ws.append([row.row_number, row.get_status_display(), row.message_text] + [row.raw_data.get(c) for c in raw_cols])
    for col, width in zip("ABC", (8, 32, 70)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def sample_workbook():
    """Documented sample format (synthetic data only)."""
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily Attendance"
    head = ["Employee ID", "Employee Name", "Attendance Date", "In Time", "Out Time", "Device"]
    ws.append(head)
    today = timezone.localdate()
    base = today - datetime.timedelta(days=today.weekday() + 7)  # Monday of last week
    sample = [("EMP1001", "Asha Demo", "09:24", "18:41"), ("EMP1002", "Rahul Sample", "09:52", "18:35"),
              ("EMP1003", "Neha Example", "10:07", "18:30"), ("EMP1004", "Vikram Test", "09:30", "")]
    for i in range(3):
        d = base + datetime.timedelta(days=i)
        for eid, name, tin, tout in sample:
            ws.append([eid, name, d.strftime("%d/%m/%Y"), tin, tout, "HO-Main-Gate"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for col in "ABCDEF":
        ws.column_dimensions[col].width = 18
    ws2 = wb.create_sheet("Punch Log")
    ws2.append(["Employee ID", "Punch Timestamp", "Punch Type", "Device"])
    for eid, _, tin, tout in sample[:2]:
        ws2.append([eid, f"{base:%d/%m/%Y} {tin}", "IN", "HO-Main-Gate"])
        ws2.append([eid, f"{base:%d/%m/%Y} 13:02", "OUT", "HO-Main-Gate"])
        ws2.append([eid, f"{base:%d/%m/%Y} 13:41", "IN", "HO-Main-Gate"])
        ws2.append([eid, f"{base:%d/%m/%Y} {tout}", "OUT", "HO-Main-Gate"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    ws3 = wb.create_sheet("Instructions")
    for line in [
        "SAMPLE FORMAT — synthetic data only. Replace with your attendance export.",
        "Daily layout: one row per employee per day. Required: Employee ID, Attendance Date, and In Time and/or Out Time.",
        "Punch-log layout: one row per punch. Required: Employee ID, Punch Timestamp. Punch Type (IN/OUT) recommended.",
        "Dates: DD/MM/YYYY (or YYYY-MM-DD, or Excel date cells). Times: HH:MM 24-hour, or h:mm AM/PM, or Excel time cells.",
        "Duplicate rule: Employee ID + timestamp (to the second) + punch type already imported → skipped.",
        "Column names can differ: you map them to system fields during upload and can save the mapping as a template.",
        "The actual vendor/biometric export format has not been supplied; confirm its layout before production use.",
    ]:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
