import datetime
import io
from unittest import mock

import openpyxl
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from apps.attendance.models import AttendanceAdjustment, AttendancePunch, AttendanceRecord
from apps.imports import services
from apps.imports.models import AttendanceImportBatch

from .utils import PASSWORD, make_employee, setup_reference

D = datetime.date
DAILY_MAP = {"employee_id": "Emp Code", "employee_name": "Name", "attendance_date": "Date", "in_time": "In", "out_time": "Out"}


def workbook(rows, header=("Emp Code", "Name", "Date", "In", "Out"), name="att.xlsx", sheet="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class ImportTests(TestCase):
    def setUp(self):
        setup_reference()
        self.admin = make_employee("ADM1", role="ADMIN")
        self.e1 = make_employee("E1001")
        self.e2 = make_employee("E1002")

    def run_import(self, upload, mapping=None, mode="DAILY", confirm=True):
        b = services.create_batch(upload, self.admin.user)
        services.detect_columns(b)
        b.format_mode, b.mapping = mode, mapping or DAILY_MAP
        b.save()
        services.validate_batch(b, self.admin.user)
        if confirm:
            b = services.import_batch(b, self.admin.user)
        b.refresh_from_db()
        return b

    def statuses(self, b):
        return list(b.rows.order_by("row_number").values_list("status", flat=True))

    def test_valid_workbook_imports_and_recalculates(self):
        b = self.run_import(workbook([("E1001", "E1001 Test", "01/09/2026", "09:35", "18:40"),
                                      ("e1002", "E1002 Test", "01/09/2026", "09:28", "18:31")]))
        self.assertEqual(b.status, "COMPLETED")
        self.assertEqual((b.imported_rows, b.punches_created), (2, 4))
        r = AttendanceRecord.objects.get(employee=self.e1, date=D(2026, 9, 1))
        self.assertEqual((r.late_minutes, r.status), (5, "PRESENT"))

    def test_invalid_workbook_rejected_at_upload(self):
        bad = SimpleUploadedFile("att.xlsx", b"PK\x03\x04 not really a zip", content_type="application/octet-stream")
        with self.assertRaises(ValidationError):
            services.create_batch(bad, self.admin.user)
        with self.assertRaises(ValidationError):
            services.create_batch(SimpleUploadedFile("att.csv", b"a,b"), self.admin.user)

    def test_missing_columns_mapping_rejected(self):
        b = services.create_batch(workbook([("E1001", "x", "01/09/2026", "09:30", "18:30")]), self.admin.user)
        services.detect_columns(b)
        b.format_mode, b.mapping = "DAILY", {"employee_id": "Emp Code"}
        b.save()
        with self.assertRaises(ValidationError):
            services.validate_batch(b, self.admin.user)
        b.mapping = {**DAILY_MAP, "in_time": "Login Time"}  # column not present
        b.save()
        with self.assertRaises(ValidationError):
            services.validate_batch(b, self.admin.user)

    def test_row_level_errors(self):
        b = self.run_import(workbook([
            ("E9999", "Ghost", "01/09/2026", "09:30", "18:30"),     # unknown employee
            ("E1001", "", "31/02/2026", "09:30", "18:30"),          # invalid date
            ("E1001", "", "02/09/2026", "9h30", "18:30"),           # invalid time
            ("", "", "02/09/2026", "09:30", "18:30"),               # missing ID
            ("E1001", "", "03/09/2026", "18:30", "09:30"),          # out before in
            ("E1001", "", "04/09/2026", "09:30", "18:30"),          # valid
        ]))
        self.assertEqual(self.statuses(b), ["ERROR", "ERROR", "ERROR", "ERROR", "ERROR", "IMPORTED"])
        self.assertEqual(b.status, "PARTIAL")
        errors = list(b.rows.filter(status="ERROR").order_by("row_number"))
        self.assertIn("Unknown employee ID", errors[0].message_text)
        self.assertIn("not a recognised date", errors[1].message_text)
        self.assertIn("not a valid time", errors[2].message_text)
        report = openpyxl.load_workbook(services.error_report_workbook(b))
        self.assertEqual(report.active.max_row, 6)  # header + 5 problem rows

    def test_duplicate_in_file_and_reimport_of_same_workbook(self):
        rows = [("E1001", "", "01/09/2026", "09:30", "18:30"), ("E1001", "", "01/09/2026", "09:30", "18:30")]
        content = workbook(rows).read()
        b = self.run_import(SimpleUploadedFile("att.xlsx", content))
        self.assertEqual(self.statuses(b), ["IMPORTED", "DUPLICATE"])
        again = self.run_import(SimpleUploadedFile("att.xlsx", content))  # byte-identical re-upload
        self.assertEqual(self.statuses(again), ["DUPLICATE", "DUPLICATE"])
        self.assertEqual(again.same_file_batch_id, b.pk)
        self.assertEqual(AttendancePunch.objects.filter(employee=self.e1).count(), 2)

    def test_multiple_punches_per_employee_per_day_punch_log(self):
        up = workbook([("E1001", "01/09/2026 09:31", "IN"), ("E1001", "01/09/2026 13:02", "OUT"),
                       ("E1001", "01/09/2026 13:40", "IN"), ("E1001", "01/09/2026 18:44", "OUT")],
                      header=("Emp Code", "Punch Time", "Direction"))
        b = self.run_import(up, mapping={"employee_id": "Emp Code", "punch_timestamp": "Punch Time", "punch_type": "Direction"}, mode="PUNCH_LOG")
        self.assertEqual(b.punches_created, 4)
        r = AttendanceRecord.objects.get(employee=self.e1, date=D(2026, 9, 1))
        self.assertEqual((r.punch_count, r.late_minutes, r.worked_minutes), (4, 1, 553))

    def test_partial_failure_rolls_back_everything(self):
        b = services.create_batch(workbook([("E1001", "", "01/09/2026", "09:30", "18:30"), ("E1002", "", "01/09/2026", "09:30", "18:30")]), self.admin.user)
        services.detect_columns(b)
        b.format_mode, b.mapping = "DAILY", DAILY_MAP
        b.save()
        services.validate_batch(b, self.admin.user)
        with mock.patch("apps.imports.services.recalculate_month", side_effect=RuntimeError("simulated crash")):
            b = services.import_batch(b, self.admin.user)
        self.assertEqual(b.status, "FAILED")
        self.assertIn("simulated crash", b.failure_message)
        self.assertEqual(AttendancePunch.objects.count(), 0)
        self.assertFalse(b.rows.filter(status="IMPORTED").exists())

    def test_corrected_source_record_requires_reconciliation(self):
        self.run_import(workbook([("E1001", "", "01/09/2026", "10:20", "18:30")], name="v1.xlsx"))
        self.assertTrue(AttendanceRecord.objects.get(employee=self.e1, date=D(2026, 9, 1)).half_day_flag)
        b2 = self.run_import(workbook([("E1001", "", "01/09/2026", "09:29", "18:30")], name="v2.xlsx"))
        row = b2.rows.get()
        self.assertEqual(row.status, "CONFLICT")
        self.assertEqual(b2.status, "PARTIAL")
        services.reconcile_row(row, self.admin.user, accept=True, note="Vendor corrected file")
        adj = AttendanceAdjustment.objects.get(import_row=row)
        self.assertEqual(adj.source, "IMPORT_RECON")
        self.assertEqual(AttendancePunch.objects.filter(employee=self.e1, is_void=False).count(), 2)  # originals kept, nothing added
        r = AttendanceRecord.objects.get(employee=self.e1, date=D(2026, 9, 1))
        self.assertFalse(r.half_day_flag)
        with self.assertRaises(ValidationError):
            services.reconcile_row(row, self.admin.user, accept=True)  # already reconciled

    def test_retry_rejected_rows_after_employee_created(self):
        b = self.run_import(workbook([("E1003", "", "01/09/2026", "09:30", "18:30")]))
        self.assertEqual(b.error_rows, 1)
        make_employee("E1003")
        self.assertEqual(services.retry_rejected_rows(b, self.admin.user), 1)
        b.refresh_from_db()
        self.assertEqual((b.status, b.imported_rows, b.error_rows), ("COMPLETED", 1, 0))

    def test_rollback_voids_and_allows_reimport(self):
        up = [("E1001", "", "01/09/2026", "09:30", "18:30")]
        b = self.run_import(workbook(up))
        services.rollback_batch(b, self.admin.user, "Wrong file")
        self.assertEqual(AttendancePunch.objects.filter(is_void=True).count(), 2)
        self.assertEqual(AttendanceRecord.objects.get(employee=self.e1, date=D(2026, 9, 1)).status, "NOT_IMPORTED")
        b2 = self.run_import(workbook(up))
        self.assertEqual(b2.imported_rows, 1)

    def test_admin_ui_end_to_end(self):
        self.client.login(username="ADM1", password=PASSWORD)
        r = self.client.post(reverse("imports:upload"), {"file": workbook([("E1001", "E1001 Test", "01/09/2026", "09:30", "18:30")])})
        batch = AttendanceImportBatch.objects.latest("pk")
        self.assertRedirects(r, reverse("imports:mapping", args=[batch.pk]))
        r = self.client.get(reverse("imports:mapping", args=[batch.pk]))
        self.assertContains(r, "Emp Code")
        data = {"sheet_name": "Sheet1", "header_row": 1, "format_mode": "DAILY", "dayfirst": "on",
                **{f"map_{k}": v for k, v in DAILY_MAP.items()}}
        r = self.client.post(reverse("imports:mapping", args=[batch.pk]), data)
        self.assertRedirects(r, reverse("imports:preview", args=[batch.pk]))
        self.assertEqual(self.client.get(reverse("imports:preview", args=[batch.pk])).status_code, 200)
        r = self.client.post(reverse("imports:confirm", args=[batch.pk]), {"confirm": "yes"})
        self.assertRedirects(r, reverse("imports:results", args=[batch.pk]))
        self.assertEqual(AttendanceImportBatch.objects.get(pk=batch.pk).status, "COMPLETED")
        self.assertEqual(self.client.get(reverse("imports:error_report", args=[batch.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("imports:sample")).status_code, 200)

    def test_non_admin_cannot_import(self):
        self.client.login(username="E1001", password=PASSWORD)
        r = self.client.post(reverse("imports:upload"), {"file": workbook([("E1001", "", "01/09/2026", "09:30", "18:30")])})
        self.assertEqual(r.status_code, 403)
        self.assertFalse(AttendanceImportBatch.objects.exists())
